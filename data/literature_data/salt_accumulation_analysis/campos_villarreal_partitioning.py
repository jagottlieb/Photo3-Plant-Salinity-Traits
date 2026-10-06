"""Cultivar-averaged Na+/Cl- concentrations per compartment and salt level, with compartment ratios.

Reads the digitized 'Campos-Villareal Ions' sheet (see digitize_campos_villarreal.py) and writes a
'Campos-Villareal Ion Ratios' sheet. Ratios are given both on absolute concentrations and on the
increase above the 0.08 dS/m control, which isolates where the added salt went.

Usage:
    python campos_villarreal_partitioning.py               # print tables
    python campos_villarreal_partitioning.py --write-xlsx  # also write the sheet into the workbook
"""
import argparse
from pathlib import Path

import pandas as pd

XLSX = Path(__file__).resolve().parents[1] / "Photosynthesis Stomatal Conductance Reduction and Dry Weight Data.xlsx"
SRC_SHEET = "Campos-Villareal Ions"
OUT_SHEET = "Campos-Villareal Ion Ratios"
CONTROL = 0.08
COMPARTMENTS = ["leaf", "stem", "root"]
RATIOS = [("leaf", "stem"), ("leaf", "root"), ("stem", "root")]


def compartment_means(ions: pd.DataFrame) -> pd.DataFrame:
    """Mean over the 4 cultivars, one row per ion and salt level, one column per compartment."""
    wide = ions.pivot_table(index=["ion", "salinity_dS_m"], columns="compartment",
                            values="value", aggfunc="mean")[COMPARTMENTS]
    return wide.sort_index(ascending=[False, True])


def add_ratios(wide: pd.DataFrame, basis: str) -> pd.DataFrame:
    out = wide.copy()
    for num, den in RATIOS:
        out[f"{num}/{den}"] = out[num] / out[den]
    out["leaf share"] = out["leaf"] / out[COMPARTMENTS].sum(axis=1)
    out.insert(0, "basis", basis)
    return out


def partition_table(ions: pd.DataFrame) -> pd.DataFrame:
    """Absolute and above-control tables stacked, each with a mean over the three salt treatments."""
    absolute = compartment_means(ions)
    control = absolute.xs(CONTROL, level="salinity_dS_m")
    above = absolute.drop(index=CONTROL, level="salinity_dS_m").sub(control, level="ion")
    blocks = []
    for wide, basis in [(absolute, "absolute"), (above, "above control")]:
        treated = wide.drop(index=CONTROL, level="salinity_dS_m", errors="ignore")
        pooled = treated.groupby(level="ion", sort=False).mean()
        pooled.index = pd.MultiIndex.from_product([pooled.index, ["2.5-3.5 mean"]],
                                                  names=wide.index.names)
        wide = wide.copy()
        wide.index = wide.index.set_levels(wide.index.levels[1].astype(str), level=1)
        blocks.append(add_ratios(pd.concat([wide, pooled]), basis))
    return pd.concat(blocks).reset_index()


def write_xlsx(table: pd.DataFrame) -> None:
    from openpyxl import load_workbook
    from openpyxl.utils.dataframe import dataframe_to_rows
    wb = load_workbook(XLSX)
    if OUT_SHEET in wb.sheetnames:
        del wb[OUT_SHEET]
    ws = wb.create_sheet(OUT_SHEET)
    ws.append([f"Cultivar-mean concentrations (mg/kg dry weight) and compartment ratios from "
               f"'{SRC_SHEET}'; made with data/literature_data/salt_accumulation_analysis/{Path(__file__).name}"])
    ws.append(["'above control' = treatment minus 0.08 dS/m control; "
               "'leaf share' = leaf / (leaf + stem + root)"])
    ws.append([])
    for r in dataframe_to_rows(table.round(3), index=False, header=True):
        ws.append(r)
    wb.save(XLSX)
    print(f"Wrote {len(table)} rows to '{OUT_SHEET}' in {XLSX.name}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--write-xlsx", action="store_true")
    args = ap.parse_args()
    ions = pd.read_excel(XLSX, sheet_name=SRC_SHEET, skiprows=2)
    table = partition_table(ions)
    with pd.option_context("display.width", 200, "display.max_columns", 20):
        print(table.round(2).to_string(index=False))
    if args.write_xlsx:
        write_xlsx(table)


if __name__ == "__main__":
    main()
